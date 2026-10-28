import io
import json
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from mumzworld_catalog import config as mconfig
from mumzworld_catalog import knowledge, pipeline, reference, samples
from mumzworld_catalog.fields import value
from mumzworld_catalog.normalize import gtin_check
from mumzworld_catalog.storage import Library, Ticket

LIB_KINDS = ["categories", "dropdowns", "mapping", "m2_template", "ns_template", "ns_master", "review_sheet", "sop"]


@pytest.fixture(autouse=True)
def mz_data(tmp_path, monkeypatch):
    monkeypatch.setattr(mconfig, "DATA_DIR", tmp_path / "mz")
    return tmp_path / "mz"


@pytest.fixture
def files(tmp_path):
    return samples.write_samples(tmp_path / "samples")


@pytest.fixture
def library(files):
    lib = Library()
    for k in LIB_KINDS:
        lib.add(k, files[k], files[k].name)
    return lib


@pytest.fixture
def ticket(library, files):
    t = Ticket.create("MW-1", "Totsy", "UAE", "2026-10-27")
    t.add_file("vendor_input", files["vendor_input"], files["vendor_input"].name)
    t.add_file("images", files["images"], files["images"].name)
    return t


def by_sku(t, sku, n=0):
    c = t.catalog()
    return [c["products"][k] for k in c["order"] if value(c["products"][k], "sku") == sku][n]


def test_gtin_check():
    assert gtin_check("5012345678900") == ""
    assert gtin_check("4006381333931") == ""
    assert "check digit" in gtin_check("5012345678901")
    assert "digits" in gtin_check("12345")


def test_reference_parsing_and_setup_check(library):
    refs = reference.Refs.from_library()
    assert len(refs.categories) == 7
    assert refs.category("Baby Gear > Strollers > Lightweight Strollers")["id"] == "101"
    assert refs.approved_values("Colour")[:2] == ["Black", "Grey"]
    assert refs.option_code("Colour", "Grey") == "c2"
    assert refs.templates["m2"]["dropdowns"]["color"][0] == "Black"
    assert refs.templates["ns"]["format"] == "csv"
    assert refs.ns_master["columns"]["sku"] == "Item Name/Number"
    assert not [c for c in reference.setup_check(refs) if c["severity"] == "blocker"]


def test_setup_check_reports_conflicts(tmp_path, library):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mappings"
    ws.append(["Target System", "Target Field", "Internal Field", "Rule", "Approved Values List"])
    ws.append(["M2", "not_a_column", "sku", "copy", ""])
    ws.append(["M2", "color", "attr:Nope", "copy", "Missing List"])
    ws.append(["NS", "Brand", "brand", "magic", ""])
    wb.save(tmp_path / "bad.xlsx")
    library.add("mapping", tmp_path / "bad.xlsx", "bad.xlsx")
    details = " ".join(c["detail"] for c in reference.setup_check(reference.Refs.from_library())
                       if c["severity"] == "blocker")
    assert "not_a_column" in details and "Missing List" in details and "unknown rule 'magic'" in details
    assert "attr:Nope" in details


def test_review_sheet_lessons_need_approval(library):
    refs = reference.Refs.from_library()
    out = knowledge.analyse_reviews(refs.reviews, ["Colour", "Age Group", "Material"])
    by_field = {x["field"]: x for x in out["lessons"]}
    assert by_field["attr:Colour"]["classification"] == "general"
    assert by_field["title"]["classification"] == "one_off"
    assert knowledge.suggest_from_lessons(out["lessons"], "test") == 1
    r = knowledge.rules()[0]
    assert r["status"] == "suggested"
    assert knowledge.applicable("attr:Colour", "Gray") is None  # not applied until approved
    with pytest.raises(ValueError):
        knowledge.decide(r["id"], "approved", "")
    knowledge.decide(r["id"], "approved", "Lead")
    assert knowledge.applicable("attr:Colour", "gray")["id"] == r["id"]


def test_conflicting_review_examples_are_not_rules():
    recs = [{"field": "Colour", "original": "Gray", "corrected": "Grey", "status": "Approved", "file": "a", "row": 2},
            {"field": "Colour", "original": "Gray", "corrected": "Silver", "status": "Approved", "file": "b",
             "row": 3, "version": "v2"}]
    out = knowledge.analyse_reviews(recs, ["Colour"])
    assert out["lessons"][0]["classification"] == "conflict"
    assert knowledge.suggest_from_lessons(out["lessons"], "x") == 0


def test_process_end_to_end(ticket):
    inv = pipeline.file_inventory(ticket)
    assert inv["vendor"][0]["records"] == 5 and inv["vendor"][0]["skus"] == 4
    assert sorted(f.name for f in ticket.path("images").iterdir()) == [
        "SIP-01.png", "TOT-200_1.png", "TOT-200_2.png", "TOT-200alt.png", "evil.png", "xxCUD-7yy.png"]
    s = pipeline.process(ticket, "tester")
    assert s["records"] == 5 and s["excluded"] == 1 and s["held"] == 1
    tot = by_sku(ticket, "TOT-200")
    f = tot["fields"]
    assert f["category"]["value"] == "Baby Gear > Strollers > Lightweight Strollers"
    assert f["category"]["class"] == "judgment" and f["category"]["status"] == "review"
    assert f["attr:Colour"]["value"] == "Grey" and f["attr:Colour"]["class"] == "mapped"  # ValueMap Charcoal
    assert f["colour"]["value"] == "Charcoal"  # vendor value kept
    assert f["title"]["value"] == "Totsy Breeze Compact Stroller - Charcoal"
    assert f["barcode"]["value"] == "5012345678917" and f["barcode"]["original"] == "5012345678917"
    assert sorted((i["file"], i["confidence"]) for i in tot["images"]) == [
        ("TOT-200_1.png", "confirmed"), ("TOT-200_2.png", "confirmed"), ("TOT-200alt.png", "uncertain")]
    assert tot["targets"]["m2"]["categories"]["value"] == "Baby Gear/Strollers/Lightweight Strollers"
    assert tot["targets"]["ns"]["Class"]["value"] == "Strollers"
    assert tot["targets"]["ns"]["Category ID"]["value"] == "101"
    old = by_sku(ticket, "TOT-100")
    assert old["match"]["status"] == "existing" and not old["eligible"]
    sip = by_sku(ticket, "SIP-01")
    assert sip["match"]["status"] == "possible_match" and not sip["eligible"]
    cud = by_sku(ticket, "CUD-7")
    assert cud["images"] == []  # 'xxCUD-7yy.png': too short an identifier to guess from
    assert "xxCUD-7yy.png" in ticket.catalog()["images"]["unassociated"]
    val = ticket.load("validation")
    codes = {(i["code"], i["sku"]) for i in val["issues"]}
    assert ("DUPLICATE_SKU", "CUD-7") in codes
    assert ("UNRESOLVED_MATCH", "SIP-01") in codes
    assert ("BARCODE_INVALID", "CUD-7") in codes
    qs = ticket.load("queries")["queries"]
    assert any(q["sku"] == "CUD-7" and q["field"] == "barcode" for q in qs)
    assert not any(q["sku"] == "TOT-200" for q in qs)  # nothing missing for a complete product
    assert ticket.meta["state"] == "awaiting_vendor"


def test_possible_match_decision_and_response_merge(ticket, files):
    pipeline.process(ticket, "tester")
    sip = by_sku(ticket, "SIP-01")
    with pytest.raises(pipeline.WorkflowError):
        pipeline.decide_match(ticket, sip["key"], "new", "Rev", "")  # reason required
    pipeline.decide_match(ticket, sip["key"], "new", "Rev", "Vendor confirmed a different product")
    sip = by_sku(ticket, "SIP-01")
    assert sip["eligible"]
    assert sip["fields"]["attr:Material"]["class"] == "vendor_query"  # mandatory for Feeding, missing
    qs = {q["field"]: q for q in ticket.load("queries")["queries"] if q["key"] == sip["key"]}
    assert qs["attr:Material"]["status"] == "open" and qs["attr:Colour"]["status"] == "open"  # 'Clear' unsupported
    r = pipeline.import_response(ticket, files["response"], files["response"].name, "tester")
    assert r["fields_updated"] == 2 and r["unmatched_skus"] == 1 and r["queries_resolved"] == 2
    sip = by_sku(ticket, "SIP-01")
    assert sip["fields"]["attr:Material"]["value"] == "Glass"
    assert sip["fields"]["colour"]["value"] == "White"
    assert sip["fields"]["colour"]["history"][0]["value"] == "Clear"
    assert sip["original"]["Colour"] == "Clear"  # original input preserved
    assert sip["fields"]["title"]["value"].endswith("- White")
    assert sip["fields"]["brand"]["value"] == "Sippy"  # unrelated fields untouched
    # a second response with a different colour is a conflict, not an overwrite
    wb = Workbook()
    wb.active.append(["SKU", "Colour"])
    wb.active.append(["SIP-01", "Black"])
    wb.save(files["response"].parent / "second.xlsx")
    r2 = pipeline.import_response(ticket, files["response"].parent / "second.xlsx", "second.xlsx", "tester")
    assert len(r2["conflicts"]) == 1 and r2["fields_updated"] == 0
    sip = by_sku(ticket, "SIP-01")
    assert sip["fields"]["colour"]["value"] == "White" and sip["fields"]["colour"]["status"] == "blocker"
    assert any(i["code"] == "RESPONSE_CONFLICT" and i["key"] == sip["key"] for i in ticket.load("validation")["issues"])
    assert len(ticket.meta["responses"]) == 2


def test_query_sheet_round_trip(ticket, tmp_path):
    pipeline.process(ticket, "tester")
    from mumzworld_catalog import exports
    refs = reference.Refs.for_ticket(ticket)
    path = tmp_path / "q.xlsx"
    exports.query_sheet(path, ticket.load("queries"), refs, ticket.meta)
    wb = load_workbook(path)
    ws = wb["Queries"]
    headers = [c.value for c in ws[1]]
    assert headers[-1] == "Vendor Answer" and "By SKU" in wb.sheetnames
    # answer the description query for the first CUD-7 row via the query id
    for row in ws.iter_rows(min_row=2):
        if row[4].value == "description":
            row[11].value = "A big soft bear."
            qid = row[0].value
    wb.save(path)
    r = pipeline.import_response(ticket, path, "answered.xlsx", "tester")
    assert r["fields_updated"] == 1
    q = next(q for q in ticket.load("queries")["queries"] if q["id"] == qid)
    assert q["status"] == "resolved" and q["answers"][0]["value"] == "A big soft bear."


def test_review_approval_and_export(ticket):
    pipeline.process(ticket, "tester")
    tot = by_sku(ticket, "TOT-200")
    with pytest.raises(pipeline.WorkflowError):
        pipeline.export(ticket, "Rev")  # nothing approved yet
    pipeline.edit_field(ticket, tot["key"], "title", "Totsy Breeze Compact Stroller, Grey", "Rev", "approved colour")
    tot = by_sku(ticket, "TOT-200")
    assert tot["fields"]["title"]["locked"] and tot["fields"]["title"]["class"] == "manual"
    assert tot["targets"]["m2"]["name"]["value"] == "Totsy Breeze Compact Stroller, Grey"
    pipeline.accept_field(ticket, tot["key"], "category", "Rev")
    cud = by_sku(ticket, "CUD-7")
    assert pipeline.review_product(ticket, [tot["key"], cud["key"]], "approved", "Rev") == 1  # CUD-7 blocked
    ap = pipeline.approve_ticket(ticket, "Rev")
    assert ap["products"] == 1 and ticket.meta["state"] == "approved"
    out = pipeline.export(ticket, "Rev")
    folder = ticket.path("outputs", out["folder"])
    wb = load_workbook(folder / out["files"]["m2"]["name"])
    assert wb.sheetnames == ["Upload", "Instructions"]
    rows = [[c.value for c in r] for r in wb["Upload"].iter_rows()]
    assert rows[0][:3] == ["sku", "name", "short_description"] and len(rows) == 2
    assert rows[1][0] == "TOT-200" and rows[1][7] == 499  # decimal column written as a number
    assert wb["Upload"].data_validations.dataValidation  # template drop-down kept
    ns = (folder / out["files"]["ns"]["name"]).read_text().splitlines()
    assert ns[0] == "Item Name/Number,Display Name,UPC Code,Class,Category ID,Brand,Cost"
    assert ns[1].startswith("TOT-200,") and "5012345678917" in ns[1] and len(ns) == 2
    rep = load_workbook(folder / out["files"]["reports"]["name"])
    assert {"Validation & Exceptions", "Existing-Product Exclusions", "Outstanding Queries",
            "Source-to-Target Mapping", "Review & Change History", "Potential New Rules"} <= set(rep.sheetnames)
    assert ticket.meta["state"] == "exported"
    # a later edit withdraws the approval
    pipeline.edit_field(ticket, tot["key"], "short_description", "Compact travel stroller.", "Rev")
    assert ticket.meta["state"] != "approved"
    with pytest.raises(pipeline.WorkflowError):
        pipeline.export(ticket, "Rev")


def test_reprocess_keeps_reviewer_work(ticket, files):
    pipeline.process(ticket, "tester")
    tot = by_sku(ticket, "TOT-200")
    pipeline.edit_field(ticket, tot["key"], "title", "Manual Title", "Rev")
    pipeline.review_product(ticket, [tot["key"]], "approved", "Rev")
    ticket.add_file("documents", files["sop"], "spec.txt")
    pipeline.process(ticket, "tester")
    tot2 = by_sku(ticket, "TOT-200")
    assert tot2["key"] == tot["key"]
    assert tot2["fields"]["title"]["value"] == "Manual Title" and tot2["fields"]["title"]["locked"]
    assert tot2["review"]["status"] == "approved"


def test_authorized_reviewers_and_waivers(ticket):
    pipeline.process(ticket, "tester")
    mconfig.save({"authorized_reviewers": ["Lead"]})
    cud = by_sku(ticket, "CUD-7", 1)
    issue = next(i for i in ticket.load("validation")["issues"] if i["key"] == cud["key"] and i["code"] == "PATTERN")
    with pytest.raises(pipeline.WorkflowError):
        pipeline.waive(ticket, issue["id"], "Someone", "ok")
    with pytest.raises(pipeline.WorkflowError):
        pipeline.waive(ticket, issue["id"], "Lead", "")
    pipeline.waive(ticket, issue["id"], "Lead", "Vendor confirmed internal code")
    w = next(i for i in ticket.load("validation")["issues"] if i["id"] == issue["id"])
    assert w["waived"]["by"] == "Lead"
    assert ticket.meta["waivers"][0]["reason"] == "Vendor confirmed internal code"


def test_unknown_column_and_category_ambiguity(library, tmp_path):
    wb2 = Workbook()
    s = wb2.active
    s.append(["Intro text"])  # the header is not on the first row
    s.append([])
    s.append(["SKU", "Brand", "Product Name", "Barcode", "Colour", "Price"])
    s.append(["ZZ-1", "Acme", "Mystery Gadget", 9780201379624.0, "Purple", 10])
    s.append(["ZZ-2", "Acme", "Bottle Teats Plush", "09780201379631", "Pink", 12])
    wb2.save(tmp_path / "v.xlsx")
    t = Ticket.create("MW-2", "Acme", "KSA", "2026-10-27")
    t.add_file("vendor_input", tmp_path / "v.xlsx", "v.xlsx")
    pipeline.process(t, "x")
    z1 = by_sku(t, "ZZ-1")
    assert z1["fields"]["barcode"]["value"] == "9780201379624"  # float cell read without '.0'
    assert z1["fields"]["category"]["class"] == "manual_review" and z1["fields"]["category"]["value"] is None
    assert z1["fields"]["attr:Colour"]["class"] == "manual_review"  # Purple is not approved
    assert any(q["field"] == "attr:Colour" and q["key"] == z1["key"] for q in t.load("queries")["queries"])
    z2 = by_sku(t, "ZZ-2")
    assert z2["fields"]["barcode"]["value"] == "09780201379631"  # leading zero kept from a text cell


def test_fake_assistant_is_constrained(ticket):
    calls = []

    class FakeMessages:
        def create(self, **kw):
            calls.append(kw)
            text = kw["messages"][0]["content"]
            text = text if isinstance(text, str) else text[-1]["text"]
            if "Pick the single most specific approved category" in text:
                out = {"category": "Not A Real Category", "confidence": "high", "rationale": "x",
                       "insufficient": False}
            elif "Attribute:" in text:
                out = {"values": ["Purple-ish", "Pink"], "equivalent": True, "rationale": "close"}
            else:
                out = {"title": "T", "short_description": "Short.", "long_description": "Clinically proven and safe.",
                       "bullet_points": ["Soft"], "search_terms": ["bear"], "unsupported_requests": []}
            return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=json.dumps(out))])

    client = SimpleNamespace(beta=SimpleNamespace(messages=FakeMessages()))
    pipeline.process(ticket, "x", ai_client=client)
    assert calls and calls[0]["model"] == "claude-opus-5-5"
    assert calls[0]["fallbacks"] == "default" and calls[0]["output_config"]["format"]["type"] == "json_schema"
    tot = by_sku(ticket, "TOT-200")
    ld = tot["fields"]["long_description"]
    assert ld["class"] == "judgment" and ld["status"] == "review"
    assert "clinically" in ld["claims"] and "safe" in ld["claims"]
    assert any(i["code"] == "UNSUPPORTED_CLAIM" and i["key"] == tot["key"] for i in ticket.load("validation")["issues"])


def test_ticket_package_round_trip(ticket):
    pipeline.process(ticket, "x")
    data = ticket.package()
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert "ticket.json" in names and "catalog.json" in names
    t2 = Ticket.import_package(data)
    assert t2.id != ticket.id and len(t2.catalog()["order"]) == 5


def test_web_flow(library, files):
    from mumzworld_catalog.web.app import create_app
    app = create_app()
    app.config["TESTING"] = True
    c = app.test_client()
    assert c.get("/").status_code == 200
    assert c.get("/library").status_code == 200
    assert c.get("/library/starter-mapping.xlsx").status_code == 200
    assert c.get("/library/reviews").status_code == 200
    with open(files["vendor_input"], "rb") as v, open(files["images"], "rb") as im:
        r = c.post("/tickets", data={"ticket_number": "MW-9", "vendor": "Totsy", "region": "UAE", "by": "Rev",
                                     "file_vendor_input": (v, "vendor.xlsx"), "file_images": (im, "imgs.zip")},
                   content_type="multipart/form-data")
    assert r.status_code == 302
    tid = r.headers["Location"].rstrip("/").split("/")[-1]
    assert b"Column mapping" in c.get(f"/t/{tid}").data
    assert c.post(f"/t/{tid}/process", data={"by": "Rev"}).status_code == 302
    t = Ticket(tid)
    key = by_sku(t, "TOT-200")["key"]
    for page in ("catalog", "queries", "responses", "validation", "review", "history", f"p/{key}"):
        resp = c.get(f"/t/{tid}/{page}")
        assert resp.status_code == 200, page
    assert c.get(f"/t/{tid}/catalog?filter=all").status_code == 200
    assert c.get(f"/t/{tid}/review?preview=1").status_code == 200
    for what in ("working", "queries", "reports", "package"):
        assert c.get(f"/t/{tid}/download/{what}").status_code == 200, what
    assert c.get(f"/t/{tid}/image/TOT-200_1.png").status_code == 200
    c.post(f"/t/{tid}/p/{key}/accept", data={"field": "category", "by": "Rev"})
    c.post(f"/t/{tid}/review/bulk", data={"key": [key], "status": "approved", "by": "Rev"})
    c.post(f"/t/{tid}/approve", data={"by": "Rev"})
    r = c.post(f"/t/{tid}/export", data={"by": "Rev"})
    assert Ticket(tid).meta["state"] == "exported"
    ex = Ticket(tid).meta["exports"][-1]
    assert c.get(f"/t/{tid}/outputs/{ex['folder']}/{ex['files']['m2']['name']}").status_code == 200
    with open(files["response"], "rb") as fh:
        r = c.post(f"/t/{tid}/responses", data={"by": "Rev", "file": (fh, "resp.xlsx")},
                   content_type="multipart/form-data")
    assert r.status_code == 200 and b"unmatched SKUs" in r.data
    assert c.get("/rules").status_code == 200
    assert c.post("/settings", data={"authorized_reviewers": "Rev"}).status_code == 302
    assert c.get("/settings").status_code == 200
