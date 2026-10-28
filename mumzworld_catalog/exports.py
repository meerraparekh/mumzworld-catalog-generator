"""Workbooks produced for a ticket: working catalog, vendor query sheet, the
Magento 2 and NetSuite import files (written into copies of the approved
templates) and the supporting reports."""
from __future__ import annotations

import csv
import io
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fields import CLASSES, label, value
from .reference import Refs

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEV_FILL = {"blocker": "FEE2E2", "warning": "FEF3C7", "info": "E5E7EB"}
CLASS_FILL = {"directly_sourced": "DBEAFE", "rule_based": "D1FAE5", "mapped": "D1FAE5", "judgment": "FEF3C7",
              "vendor_query": "FEE2E2", "manual_review": "FEE2E2", "manual": "EDE9FE"}


NUMBER_TYPES = {"number", "decimal", "integer", "int", "price", "float"}


class ExportError(RuntimeError):
    pass


def _number(v):
    try:
        f = float(str(v).replace(",", ""))
    except ValueError:
        return None
    return int(f) if f.is_integer() else f


def _sheet(wb, title, columns, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(columns)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for c in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, start=1):
        w = (widths or {}).get(col)
        if not w:
            longest = max([len(str(col))] + [len(str(r[i - 1]).split("\n")[0]) for r in rows[:200] if i - 1 < len(r)])
            w = max(10, min(60, longest + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


# -- working catalog ---------------------------------------------------------------------------
def working_catalog(path: Path, catalog: dict, refs: Refs, validation: dict) -> None:
    products, order = catalog["products"], catalog["order"]
    attrs = [a["attribute"] for a in (refs.mapping or {}).get("attributes", [])]
    blocked = set(validation.get("blocked_keys", []))
    main_fields = ["sku", "brand", "product_name", "barcode", "category", "title", "short_description",
                   "long_description", "bullet_points", "colour", "size", "material", "age_group", "weight",
                   "dimensions"]
    cols = ["Key", "Vendor File", "Row", "Match", "Eligible", "Readiness", "Review"] + [label(f) for f in main_fields] \
        + attrs + ["Images"]
    rows = []
    for k in order:
        p = products[k]
        rows.append([k, p["source"]["file"], p["source"]["row"], p.get("match", {}).get("label", ""),
                     "Yes" if p.get("eligible") else "No",
                     ("Blocked" if k in blocked else "Ready") if p.get("eligible") else "Excluded",
                     p.get("review", {}).get("status", "")] + [value(p, f) for f in main_fields]
                    + [value(p, f"attr:{a}") for a in attrs]
                    + [", ".join(i["file"] for i in p.get("images", []))])
    wb = Workbook()
    ws = _sheet(wb, "Working Catalog", cols, rows, first=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    detail = []
    for k in order:
        p = products[k]
        for f, d in p["fields"].items():
            detail.append([k, value(p, "sku"), label(f), f, d.get("value"), d.get("original"),
                           CLASSES.get(d.get("class"), d.get("class")), d.get("source"), d.get("transform"),
                           d.get("reason"), d.get("status"), "Yes" if d.get("locked") else "",
                           len(d.get("history", []))])
    ws = _sheet(wb, "Field Detail", ["Key", "SKU", "Field", "Field Code", "Current Value", "Original Value",
                                     "Decision Class", "Source", "Transformation", "Reason", "Status", "Locked",
                                     "Changes"], detail)
    for r in ws.iter_rows(min_row=2):
        fill = CLASS_FILL.get(next((k for k, v in CLASSES.items() if v == r[6].value), ""))
        if fill:
            r[6].fill = PatternFill("solid", fgColor=fill)
    for system, title in (("m2", "M2 Preview"), ("ns", "NS Preview")):
        t = refs.templates.get(system)
        if not t:
            continue
        names = [c["name"] for c in t["columns"]]
        _sheet(wb, title, ["Key"] + names, [[k] + [(products[k].get("targets") or {}).get(system, {}).get(n, {}).get(
            "value", "") for n in names] for k in order if products[k].get("eligible")])
    wb.save(path)


# -- vendor queries ----------------------------------------------------------------------------
def query_sheet(path: Path, queries: dict, refs: Refs, meta: dict) -> None:
    open_q = [q for q in queries.get("queries", []) if q["status"] in ("open", "partially", "conflict")]
    wb = Workbook()
    cols = ["Query ID", "SKU", "Product", "Field", "Field Code", "Current Value", "Question", "Reason",
            "Expected Format / Allowed Values", "Priority", "Status", "Vendor Answer"]
    ws = _sheet(wb, "Queries", cols, [[q["id"], q["sku"], q["product"], q["field_label"], q["field"], q.get("current"),
                                       q["question"], q["reason"], q.get("expected"), q["priority"], q["status"], ""]
                                      for q in sorted(open_q, key=lambda x: (x["sku"], x["id"]))],
                widths={"Question": 60, "Reason": 50, "Expected Format / Allowed Values": 50, "Vendor Answer": 40},
                first=True)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r[11].fill = PatternFill("solid", fgColor="FEF9C3")
    by_sku = {}
    for q in open_q:
        by_sku.setdefault((q["sku"], q["product"]), []).append(q)
    _sheet(wb, "By SKU", ["SKU", "Product", "Questions", "Highest Priority"],
           [[s, n, "\n".join(f"{i + 1}. {q['question']}" + (f" ({q['expected']})" if q.get("expected") else "")
                             for i, q in enumerate(qs)),
             "High" if any(q["priority"] == "High" for q in qs) else "Medium"] for (s, n), qs in sorted(by_sku.items())],
           widths={"Questions": 100})
    lists = sorted({a["list"] for a in (refs.mapping or {}).get("attributes", [])
                    if any(q["field"] == f"attr:{a['attribute']}" for q in open_q)})
    if lists:
        mx = max(len(refs.approved_values(lst) or []) for lst in lists)
        _sheet(wb, "Allowed Values", lists, [[(refs.approved_values(lst) or [])[i] if i < len(
            refs.approved_values(lst) or []) else "" for lst in lists] for i in range(mx)])
    info = wb.create_sheet("About")
    for line in (["Ticket", meta.get("ticket_number")], ["Vendor", meta.get("vendor")],
                 ["How to answer", "Fill the Vendor Answer column on the Queries sheet and send the file back. "
                                   "Keep the Query ID, SKU and Field Code columns unchanged."]):
        info.append(line)
    wb.save(path)


# -- final import files ------------------------------------------------------------------------
def final_output(path: Path, refs: Refs, system: str, products: list[dict]) -> dict:
    """Write approved products into a copy of the approved template.  Returns a summary."""
    t = refs.templates.get(system)
    if not t:
        raise ExportError(f"The {system.upper()} template is not available.")
    src = Path(t["path"])
    names = [c["name"] for c in t["columns"]]
    rows = [[(p.get("targets") or {}).get(system, {}).get(n, {}) for n in names] for p in products]
    if t["format"] == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        original = src.read_text("utf-8-sig", errors="replace").splitlines()
        for line in original[: t["header_row"]]:
            buf.write(line + "\r\n")
        for r in rows:
            w.writerow([c.get("value", "") for c in r])
        path.write_text(buf.getvalue(), "utf-8")
        return {"rows": len(rows), "replaced_rows": t["existing_rows"], "columns": len(names)}
    shutil.copyfile(src, path)
    try:
        wb = load_workbook(path, keep_vba=t["format"] == "xlsm")
    except Exception as e:
        raise ExportError(f"The template could not be opened for writing: {e}") from e
    ws = wb[t["sheet"]]
    # verify the header row is exactly as parsed before writing anything
    for c in t["columns"]:
        if str(ws.cell(row=t["header_row"], column=c["index"] + 1).value).strip() != c["name"]:
            raise ExportError(f"Template header '{c['name']}' is not where expected; export stopped.")
    first = t["header_row"] + 1
    for r in range(first, ws.max_row + 1):  # clear old example rows, keep formatting
        for c in t["columns"]:
            ws.cell(row=r, column=c["index"] + 1).value = None
    for i, r in enumerate(rows):
        for c, cell in zip(t["columns"], r):
            v = cell.get("value", "")
            target = ws.cell(row=first + i, column=c["index"] + 1)
            if v == "":
                target.value = None
                continue
            if cell.get("text"):
                target.number_format = "@"
                target.value = str(v)
            elif cell.get("type") in NUMBER_TYPES and _number(v) is not None:
                target.value = _number(v)
            else:
                target.value = v
    wb.save(path)
    return {"rows": len(rows), "replaced_rows": t["existing_rows"], "columns": len(names)}


# -- supporting reports ------------------------------------------------------------------------
def reports(path: Path, ticket_meta: dict, catalog: dict, refs: Refs, validation: dict, queries: dict,
            exclusions: list[dict], lessons: dict) -> None:
    products, order = catalog["products"], catalog["order"]
    wb = Workbook()
    _sheet(wb, "Validation & Exceptions",
           ["Severity", "Code", "SKU", "Field", "System", "Current Value", "Expected", "Reason", "Suggested Action",
            "Waived By", "Waiver Reason"],
           [[i["severity"], i["code"], i["sku"], label(i["field"]) if i["field"] else "", i["system"].upper(),
             i["current"], i["expected"], i["reason"], i["action"], (i.get("waived") or {}).get("by", ""),
             (i.get("waived") or {}).get("reason", "")] for i in validation.get("issues", [])], first=True,
           widths={"Reason": 60, "Suggested Action": 40})
    ex = [[e["sku"], e["barcode"], e["match_identifier"], e["ns_record"], e["ns_row"], e["match_type"], e["decision"],
           e["reason"]] for e in exclusions]
    ex += [["", "", "", "", "", x["type"], "Not merged", f"{x.get('detail', '')} ({x.get('file', '')} "
            f"{x.get('sheet', '')} row {x.get('row', '')}, SKU {x.get('sku', '')})"]
           for x in catalog.get("exceptions", [])]
    _sheet(wb, "Existing-Product Exclusions", ["SKU", "Barcode", "Match Identifier", "Matching NetSuite Record",
                                               "NS Row", "Match Type", "Decision", "Reason"], ex,
           widths={"Reason": 70})
    _sheet(wb, "Outstanding Queries", ["Query ID", "SKU", "Field", "Question", "Priority", "Status", "Answers"],
           [[q["id"], q["sku"], q["field_label"], q["question"], q["priority"], q["status"],
             "; ".join(f"{a['value']} ({a['file']})" for a in q["answers"])]
            for q in queries.get("queries", []) if q["status"] in ("open", "partially", "conflict")])
    mrows = []
    for r in (refs.mapping or {}).get("mappings", []):
        mrows.append([r["system"].upper(), r["target"], r["source"], r["rule"], r["param"], "Y" if r["required"] else "",
                      r["type"], r["list"], r["validation"], r["max_length"] or "", r["reference"], r["row"]])
    _sheet(wb, "Source-to-Target Mapping", ["System", "Target Field", "Internal Field / Source", "Rule",
                                            "Parameter", "Required", "Data Type", "Approved List", "Validation",
                                            "Max Length", "Approval Reference", "Mapping Row"], mrows)
    trows = []
    for k in order:
        p = products[k]
        if not p.get("eligible"):
            continue
        for f, d in p["fields"].items():
            if d.get("transform") or d.get("class") not in ("directly_sourced",):
                trows.append([value(p, "sku"), label(f), d.get("original"), d.get("value"),
                              CLASSES.get(d.get("class"), d.get("class")), d.get("source"), d.get("transform"),
                              d.get("reason")])
    _sheet(wb, "Transformations", ["SKU", "Field", "Source Value", "Working Value", "Decision Class", "Source",
                                   "Transformation", "Reason"], trows, widths={"Reason": 60})
    hist = [[h["at"], h.get("by", ""), h["event"]] for h in ticket_meta.get("history", [])]
    for k in order:
        p = products[k]
        for f, d in p["fields"].items():
            for h in d.get("history", []):
                hist.append([h["at"], "", f"{value(p, 'sku')} {label(f)}: was '{h.get('value')}' "
                                          f"({h.get('source')}, {CLASSES.get(h.get('class'), h.get('class'))})"])
        for h in p.get("review_log", []):
            hist.append([h["at"], h.get("by", ""), f"{value(p, 'sku')}: {h['event']}"])
    hist.sort(key=lambda x: x[0])
    _sheet(wb, "Review & Change History", ["When", "Who", "Event"], hist, widths={"Event": 100})
    _sheet(wb, "Potential New Rules", ["Classification", "Field", "From", "To", "Scope", "Examples", "Note"],
           [[x["classification"], label(x["field"]), x["from"], x["to"],
             ", ".join(f"{k2}: {v2}" for k2, v2 in x.get("scope", {}).items() if v2), len(x["examples"]),
             x.get("note", "")] for x in lessons.get("lessons", [])]
           + [["observation", "", "", "", "", n, text] for text, n in lessons.get("observations", [])])
    used = ticket_meta.get("refs_used", {})
    _sheet(wb, "References Used", ["Kind", "File", "Origin", "Version / File ID", "SHA-1"],
           [[k, e["name"], e["origin"], e["id"], e["sha1"]] for k, es in used.items() for e in es])
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        if ws.title == "Validation & Exceptions":
            for row in ws.iter_rows(min_row=2):
                fill = SEV_FILL.get(row[0].value)
                if fill:
                    row[0].fill = PatternFill("solid", fgColor=fill)
    wb.save(path)
