"""Reading the reference library: categories, dropdowns, mapping rules, templates,
the NetSuite master, review sheets and SOP text.

Column roles are recognised from header names.  What could not be recognised
is reported in the setup summary instead of being guessed, and nothing here
creates a category, dropdown value or mapping that is not in the files.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from openpyxl.utils import get_column_letter

from . import cache
from . import sheets as sh
from .fields import SOURCE_FIELDS, resolve_field_name
from .sheets import as_text, is_blank, norm
from .storage import MULTI_KINDS, REF_KINDS, REQUIRED_KINDS, Library


def _col(columns: list[str], *names: str, contains: tuple[str, ...] = ()) -> str | None:
    keys = {norm(n) for n in names}
    for c in columns:
        if norm(c) in keys:
            return c
    for c in columns:
        if any(x in norm(c) for x in contains):
            return c
    return None


def _yes(v) -> bool:
    return norm(v) in {"y", "yes", "true", "1", "mandatory", "required", "m", "x"}


def _split(v, seps=r"[|;\n]") -> list[str]:
    return [x.strip() for x in re.split(seps, as_text(v)) if x.strip()]


# -- categories -----------------------------------------------------------------------
_LEVEL_HINTS = ["level", "lvl", "category", "subcategory", "subsubcategory", "department", "family", "class",
                "subclass", "l1", "l2", "l3", "l4", "l5", "tier"]


def parse_categories(path: Path) -> dict:
    t = sh.read_table(path)
    cols = t["columns"]
    id_col = _col(cols, "category id", "id", "cat id", "category code", "code", "magento id", "m2 category id",
                  "entity id", "categoryid")
    path_col = _col(cols, "path", "category path", "full path", "full category path", "breadcrumb",
                    "category tree")
    kw_col = _col(cols, "keywords", "synonyms", "includes", "product types", "search terms", "tagging keywords")
    special_col = _col(cols, "special", "special category", "exception", "requires approval", "special handling")
    multi_col = _col(cols, "allow multiple", "multi category", "multiple categories allowed", "multi assign",
                     "allow multi category")
    rules_col = _col(cols, "rules", "notes", "tagging rule", "guideline", "comments")
    used = {c for c in (id_col, path_col, kw_col, special_col, multi_col, rules_col) if c}
    level_cols = [c for c in cols if c not in used and (
        any(norm(c).startswith(h) for h in _LEVEL_HINTS) or re.fullmatch(r"(sub)*(category|cat)\d*", norm(c)))]
    problems = []
    if not path_col and not level_cols:
        raise sh.SheetError("No category path or level columns were recognised (expected headers such as "
                            "'Category', 'Subcategory', 'Level 1' or 'Path').")
    sep = " > "
    if path_col:
        sample = " ".join(as_text(r["values"].get(path_col)) for r in t["records"][:50])
        for cand in (" > ", ">", " / ", "/", "|", "\\"):
            if cand in sample:
                sep = cand
                break
    cats, seen = [], set()
    for r in t["records"]:
        v = r["values"]
        if path_col and not is_blank(v.get(path_col)):
            parts = [p.strip() for p in as_text(v[path_col]).split(sep.strip() or sep) if p.strip()]
        else:
            parts = [as_text(v.get(c)).strip() for c in level_cols if not is_blank(v.get(c))]
        if not parts:
            continue
        key = " > ".join(parts)
        if key in seen:
            problems.append(f"Category '{key}' is listed more than once (row {r['row']}).")
            continue
        seen.add(key)
        cats.append({
            "key": key, "path": parts, "id": as_text(v.get(id_col)).strip() if id_col else "",
            "keywords": _split(v.get(kw_col), r"[|;,\n]") if kw_col else [],
            "special": _yes(v.get(special_col)) if special_col else False,
            "multi": _yes(v.get(multi_col)) if multi_col else False,
            "rule": as_text(v.get(rules_col)) if rules_col else "",
            "extra": {c: as_text(v.get(c)) for c in cols if c not in level_cols and c != path_col
                      and not is_blank(v.get(c))},
        })
    return {"sheet": t["sheet"], "categories": cats, "problems": problems,
            "columns": {"id": id_col, "path": path_col, "levels": level_cols, "keywords": kw_col,
                        "special": special_col, "multi": multi_col}}


# -- dropdowns ------------------------------------------------------------------------
def parse_dropdowns(path: Path) -> dict:
    """{attribute: {"values": [...], "codes": {value: code}, "source": "sheet"}}"""
    out: dict[str, dict] = {}
    problems = []
    for name, rows in sh.sheet_rows(path).items():
        if not rows:
            continue
        t = sh.table_from_rows(rows, sh.header_row_index(rows))
        cols = t["columns"]
        attr_col = _col(cols, "attribute", "attribute name", "attribute code", "field", "list", "list name",
                        "dropdown", "dropdown name", "attribute label")
        val_col = _col(cols, "value", "option", "option value", "allowed value", "approved value", "label",
                       "option label", "values", "dropdown value")
        code_col = _col(cols, "option id", "value id", "code", "option code", "internal id", "value code")
        if attr_col and val_col:
            for r in t["records"]:
                a, v = as_text(r["values"].get(attr_col)).strip(), as_text(r["values"].get(val_col)).strip()
                if not a or not v:
                    continue
                e = out.setdefault(a, {"values": [], "codes": {}, "source": f"{name} (list)"})
                if v in e["values"]:
                    continue
                e["values"].append(v)
                if code_col and not is_blank(r["values"].get(code_col)):
                    e["codes"][v] = as_text(r["values"][code_col]).strip()
        else:  # wide: each column is one attribute
            for c in cols:
                if c.startswith("Column "):
                    continue
                vals = []
                for r in t["records"]:
                    v = as_text(r["values"].get(c)).strip()
                    if v and v not in vals:
                        vals.append(v)
                if not vals:
                    continue
                if c in out:
                    problems.append(f"Attribute '{c}' appears in more than one sheet; values were combined.")
                e = out.setdefault(c, {"values": [], "codes": {}, "source": f"{name} (columns)"})
                e["values"] += [v for v in vals if v not in e["values"]]
    return {"attributes": out, "problems": problems}


# -- mapping workbook -----------------------------------------------------------------
MAPPING_RULES = {
    "copy": "Value of the internal field (or attr:<Attribute>) named in Source",
    "constant": "The text in Rule Parameter",
    "template": "Text with {field} placeholders, e.g. '{brand} {title}'",
    "lookup": "Exact match of the source value in the ValueMap named in Rule Parameter",
    "concat": "Join the '|'-separated Source fields with the separator in Rule Parameter",
    "category_path": "Chosen category path joined with the separator in Rule Parameter (default ' > ')",
    "category_level": "One level of the chosen category path (Rule Parameter = level number, 1 = top)",
    "category_id": "ID of the chosen category from the category list",
    "category_field": "Another column of the category list for the chosen category (Rule Parameter = column)",
    "attribute_code": "Option code of the mapped dropdown value (Source = attr:<Attribute>)",
    "images": "Associated image file names joined with Rule Parameter (default ',')",
    "blank": "Always left empty (column kept in the template)",
}
CONTENT_RULE_KEYS = {
    "title_pattern": "Title built from {field} placeholders, e.g. '{brand} {product_name} - {colour}'",
    "title_max_length": "Maximum title length in characters",
    "title_case": "title | sentence | as_is",
    "short_description_max_length": "Maximum short description length",
    "long_description_min_length": "Minimum long description length (shorter is a warning)",
    "bullets_max": "Maximum number of key-feature bullets",
    "bullet_max_length": "Maximum length of one bullet",
    "forbidden_terms": "'|'-separated terms that must not appear in content",
    "claim_terms": "'|'-separated claim words that need supporting evidence in vendor data",
    "images_required": "Y if every product needs at least one image",
    "min_images": "Minimum number of images per product",
    "multi_category": "Y if more than one category may be assigned (only where the category list allows it)",
}
DEFAULT_CLAIM_TERMS = ["certified", "certification", "approved", "bpa free", "bpa-free", "non-toxic", "hypoallergenic",
                       "organic", "safe", "safest", "medical", "clinically", "dermatologically", "tested",
                       "antibacterial", "anti-bacterial", "cure", "prevents", "eco-friendly", "fda", "ce marked",
                       "en 71", "en71", "astm", "iso", "guaranteed", "100%", "best", "No. 1"]


def parse_mapping(path: Path) -> dict:
    out = {"settings": {}, "mappings": [], "attributes": [], "content_rules": [], "value_maps": {},
           "problems": []}
    all_sheets = sh.sheet_rows(path)
    by_norm = {norm(n): n for n in all_sheets}

    def table(*names):
        for n in names:
            if norm(n) in by_norm:
                rows = all_sheets[by_norm[norm(n)]]
                return sh.table_from_rows(rows, sh.header_row_index(rows)) if rows else None
        return None

    t = table("Settings")
    if t:
        kc = _col(t["columns"], "setting", "key", "name") or t["columns"][0]
        vc = _col(t["columns"], "value") or (t["columns"][1] if len(t["columns"]) > 1 else None)
        for r in t["records"]:
            k = as_text(r["values"].get(kc)).strip()
            if k and vc:
                out["settings"][norm(k)] = as_text(r["values"].get(vc)).strip()
    t = table("Mappings", "Field Mappings", "Mapping")
    if not t:
        out["problems"].append("No 'Mappings' sheet: Magento 2 and NetSuite columns have no confirmed mapping.")
    else:
        c = t["columns"]
        sys_c = _col(c, "target system", "system")
        tgt_c = _col(c, "target field", "target column", "target attribute", "column")
        src_c = _col(c, "internal field", "source", "source field", "source fields", "internal field source")
        rule_c = _col(c, "rule", "transformation rule", "transformation")
        par_c = _col(c, "rule parameter", "parameter", "param")
        req_c = _col(c, "required", "mandatory")
        type_c = _col(c, "data type", "type")
        list_c = _col(c, "approved values list", "approved list", "dropdown", "values list")
        val_c = _col(c, "validation rule", "validation", "regex", "pattern")
        len_c = _col(c, "max length", "maximum length", "length")
        ref_c = _col(c, "source or approval reference", "approval reference", "reference", "approved by")
        sep_c = _col(c, "multi value separator", "separator")
        if not tgt_c or not sys_c:
            out["problems"].append("Mappings sheet needs 'Target System' and 'Target Field' columns.")
        else:
            for r in t["records"]:
                v = r["values"]
                system = norm(v.get(sys_c))
                system = {"m2": "m2", "magento": "m2", "magento2": "m2", "ns": "ns", "netsuite": "ns"}.get(system)
                target = as_text(v.get(tgt_c)).strip()
                if not target:
                    continue
                if not system:
                    out["problems"].append(f"Mappings row {r['row']}: Target System must be M2 or NS.")
                    continue
                rule = norm(v.get(rule_c)) if rule_c else ""
                rule = {"categorypath": "category_path", "categorylevel": "category_level",
                        "categoryid": "category_id", "categoryfield": "category_field",
                        "attributecode": "attribute_code"}.get(rule, rule) or ("copy" if src_c and not is_blank(
                            v.get(src_c)) else "")
                out["mappings"].append({
                    "row": r["row"], "system": system, "target": target,
                    "source": as_text(v.get(src_c)).strip() if src_c else "", "rule": rule,
                    "param": as_text(v.get(par_c)) if par_c else "",
                    "required": _yes(v.get(req_c)) if req_c else False,
                    "type": norm(v.get(type_c)) if type_c else "",
                    "list": as_text(v.get(list_c)).strip() if list_c else "",
                    "validation": as_text(v.get(val_c)).strip() if val_c else "",
                    "max_length": int(float(v[len_c])) if len_c and _num(v.get(len_c)) else None,
                    "reference": as_text(v.get(ref_c)).strip() if ref_c else "",
                    "separator": as_text(v.get(sep_c)) if sep_c else "",
                })
    t = table("Attributes", "Attribute Rules", "Attribute Mapping")
    if t:
        c = t["columns"]
        a_c = _col(c, "attribute", "attribute name")
        f_c = _col(c, "internal field", "source field", "source", "vendor field")
        cat_c = _col(c, "applies to categories", "categories", "category", "applies to")
        m_c = _col(c, "mandatory", "required")
        l_c = _col(c, "approved values list", "dropdown", "list", "approved list")
        s_c = _col(c, "multi value separator", "separator", "multi value")
        fmt_c = _col(c, "expected format", "format", "vendor query format")
        for r in t["records"]:
            v = r["values"]
            name = as_text(v.get(a_c)).strip() if a_c else ""
            if not name:
                continue
            out["attributes"].append({
                "row": r["row"], "attribute": name,
                "source": as_text(v.get(f_c)).strip() if f_c else "",
                "categories": _split(v.get(cat_c)) if cat_c else [],
                "mandatory": _yes(v.get(m_c)) if m_c else False,
                "list": (as_text(v.get(l_c)).strip() if l_c else "") or name,
                "separator": as_text(v.get(s_c)) if s_c else "",
                "format": as_text(v.get(fmt_c)).strip() if fmt_c else "",
            })
    t = table("ContentRules", "Content Rules", "Content")
    if t:
        c = t["columns"]
        k_c = _col(c, "rule", "key", "setting") or c[0]
        v_c = _col(c, "value") or (c[1] if len(c) > 1 else None)
        cat_c = _col(c, "applies to category", "category", "applies to")
        for r in t["records"]:
            k = norm(r["values"].get(k_c))
            key = next((x for x in CONTENT_RULE_KEYS if norm(x) == k), None)
            if not key:
                if k:
                    out["problems"].append(f"ContentRules row {r['row']}: unknown rule "
                                           f"'{as_text(r['values'].get(k_c))}'.")
                continue
            out["content_rules"].append({"row": r["row"], "rule": key,
                                         "value": as_text(r["values"].get(v_c)).strip() if v_c else "",
                                         "category": as_text(r["values"].get(cat_c)).strip() if cat_c else ""})
    t = table("ValueMap", "Value Map", "ValueMaps")
    if t:
        c = t["columns"]
        m_c = _col(c, "map name", "map", "list", "attribute")
        s_c = _col(c, "source value", "vendor value", "from")
        d_c = _col(c, "target value", "approved value", "to")
        if m_c and s_c and d_c:
            for r in t["records"]:
                v = r["values"]
                m, s, d = (as_text(v.get(x)).strip() for x in (m_c, s_c, d_c))
                if m and s:
                    out["value_maps"].setdefault(m, {})[norm(s)] = {"source": s, "target": d, "row": r["row"]}
    return out


def _num(v) -> bool:
    try:
        float(as_text(v))
        return True
    except ValueError:
        return False


# -- templates ------------------------------------------------------------------------
def parse_template(path: Path, sheet: str = "", header_row: int | None = None) -> dict:
    path = Path(path)
    all_sheets = sh.sheet_rows(path)
    if not all_sheets:
        raise sh.SheetError("The template has no sheets.")
    if sheet and sheet not in all_sheets:
        raise sh.SheetError(f"Template sheet '{sheet}' not found. Sheets: {', '.join(all_sheets)}")
    if not sheet:
        named = [s for s in all_sheets if re.search(r"upload|import|template|product|item", s, re.I) and all_sheets[s]]
        sheet = (named or [s for s in all_sheets if all_sheets[s]] or list(all_sheets))[0]
    rows = all_sheets[sheet]
    h = (header_row - 1) if header_row else sh.header_row_index(rows)
    headers = rows[h] if rows else []
    columns = []
    for i, v in enumerate(headers):
        if not is_blank(v):
            columns.append({"name": str(v).strip(), "index": i, "letter": get_column_letter(i + 1)})
    if not columns:
        raise sh.SheetError(f"No header row found on sheet '{sheet}'.")
    required = set()
    for marker_row in (h - 1, h + 1):  # a "Required/Optional" row next to the headers
        if 0 <= marker_row < len(rows):
            r = rows[marker_row]
            marks = {c["name"]: norm(r[c["index"]]) if c["index"] < len(r) else "" for c in columns}
            if sum(1 for m in marks.values() if m in {"required", "mandatory", "optional"}) >= max(2, len(columns) // 4):
                required |= {n for n, m in marks.items() if m in {"required", "mandatory"}}
    required |= {c["name"] for c in columns if c["name"].endswith("*")}
    data_rows = sum(1 for r in rows[h + 1:] if not all(is_blank(x) for x in r))
    dvs = sh.data_validations(path).get(sheet, {})
    dropdowns = {c["name"]: dvs[c["letter"]] for c in columns if c["letter"] in dvs}
    return {"file": path.name, "format": path.suffix.lower().lstrip("."), "sheet": sheet, "header_row": h + 1,
            "columns": columns, "required": sorted(required), "dropdowns": dropdowns,
            "existing_rows": data_rows, "other_sheets": [s for s in all_sheets if s != sheet]}


# -- NetSuite master ------------------------------------------------------------------
NS_SKU_NAMES = SOURCE_FIELDS["sku"] + ["item name/number", "item name number", "itemid", "item id", "item number",
                                       "name"]


NS_FIELDS = ["sku", "internal_id", "barcode", "vendor_sku", "name", "brand", "category", "status"]


def _ns_columns(c: list[str]) -> dict | None:
    sku_c = _col(c, *SOURCE_FIELDS["sku"]) or _col(c, "item name/number", "item name number", "itemid", "item id",
                                                     "item number", "name")
    if not sku_c:
        return None
    cols = {
        "sku": sku_c,
        "internal_id": _col(c, "internal id", "internalid", "id", "ns id", "netsuite id"),
        "barcode": _col(c, "upc code", "upc", "barcode", "ean", "gtin", "ean code", "ean/upc", "upc/ean"),
        "vendor_sku": _col(c, "vendor sku", "vendor code", "vendor item code", "vendor name/code", "supplier sku",
                           "mpn", "manufacturer part number", "vendor part number"),
        "name": _col(c, "display name", "display name/code", "item display name", "description", "product name",
                     "sales description", "title"),
        "brand": _col(c, "brand", "manufacturer"),
        "category": _col(c, "category", "class", "item category", "product category", "department"),
        "status": _col(c, "status", "inactive", "is inactive", "item status"),
    }
    if cols["internal_id"] == sku_c:
        cols["internal_id"] = None
    return cols


def parse_ns_master(path: Path, progress=None) -> dict:
    """Stream every sheet that has a SKU column, keeping only the columns used for
    matching.  Records are compact rows ``[sheet index, row, *NS_FIELDS]``."""
    names, heads = sh.pick_sheets(path, NS_SKU_NAMES)
    sheets, rows, all_cols = [], [], {}
    for name in names:
        head = heads[name]
        h = sh.header_row_index(head, NS_SKU_NAMES)
        c = sh.table_from_rows(head[: h + 1], h)["columns"]
        cols = _ns_columns(c)
        if not cols:
            continue
        idx = [c.index(cols[f]) if cols[f] else None for f in NS_FIELDS]
        si = len(sheets)
        sheets.append({"name": name, "columns": cols, "header_row": h + 1})
        all_cols[name] = c
        if Path(path).suffix.lower() == ".csv":
            stream = ((n, dict(enumerate(r))) for n, r in enumerate(sh.iter_rows(path), start=1))
        else:
            stream = sh.fast_rows(path, name, {i for i in idx if i is not None})
        for n, cells in stream:
            if n <= h + 1:
                continue
            vals = [as_text(cells.get(i)).strip() if i is not None else "" for i in idx]
            if not any(vals):
                continue
            rows.append([si, n] + vals)
            if progress and len(rows) % 5000 == 0:
                progress(len(rows))
    if not sheets:
        raise sh.SheetError("No SKU column was recognised in the NetSuite master (expected e.g. 'SKU', "
                            "'Item Name/Number' or 'Name').")
    return {"sheets": sheets, "sheet": ", ".join(x["name"] for x in sheets), "columns": sheets[0]["columns"],
            "all_columns": all_cols, "fields": NS_FIELDS, "rows": rows}


def ns_record(master: dict, row: list) -> dict:
    """One compact master row as a dict (sheet, row and the NS_FIELDS)."""
    rec = dict(zip(master["fields"], row[2:]))
    rec["sheet"] = master["sheets"][row[0]]["name"]
    rec["row"] = row[1]
    return rec


# -- review sheets --------------------------------------------------------------------
def parse_review_sheet(path: Path) -> dict:
    out = {"records": [], "problems": [], "sheets": []}
    for name, rows in sh.sheet_rows(path).items():
        if not rows:
            continue
        t = sh.table_from_rows(rows, sh.header_row_index(rows))
        c = t["columns"]
        cols = {
            "sku": _col(c, *SOURCE_FIELDS["sku"]),
            "field": _col(c, "field", "attribute", "field reviewed", "attribute reviewed", "column"),
            "original": _col(c, "original value", "original", "before", "old value", "proposed value",
                             "tool value", "system value", "submitted value", "initial value"),
            "corrected": _col(c, "corrected value", "corrected", "after", "new value", "approved value",
                              "final value", "correct value", "reviewed value"),
            "comment": _col(c, "reviewer comments", "reviewer comment", "comments", "comment", "remarks", "notes",
                            "feedback"),
            "reason": _col(c, "reason", "reason for correction", "correction reason", "root cause"),
            "status": _col(c, "approval status", "status", "approved", "review status"),
            "category": _col(c, "category", "product type", "category path"),
            "brand": _col(c, "brand"),
            "version": _col(c, "rule version", "sop version", "version"),
            "date": _col(c, "date", "review date", "reviewed on"),
            "reviewer": _col(c, "reviewer", "reviewed by", "approved by"),
        }
        if not (cols["field"] and cols["corrected"]):
            out["problems"].append(f"Sheet '{name}': no 'Field' and 'Corrected value' columns; not read as "
                                   "review decisions.")
            continue
        out["sheets"].append(name)
        for r in t["records"]:
            v = r["values"]
            rec = {k: (as_text(v.get(col)).strip() if col else "") for k, col in cols.items()}
            if not rec["field"]:
                continue
            rec["row"], rec["sheet"] = r["row"], name
            out["records"].append(rec)
    return out


# -- SOP / documents ------------------------------------------------------------------
def extract_text(path: Path, limit: int = 200_000) -> tuple[str, str]:
    """(text, note).  The note says when text could not be extracted."""
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext in (".txt", ".md", ".csv"):
            return p.read_text("utf-8", errors="replace")[:limit], ""
        if ext == ".docx":
            with zipfile.ZipFile(p) as zf:
                xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
            xml = re.sub(r"</w:p>", "\n", xml)
            return re.sub(r"<[^>]+>", "", xml)[:limit], ""
        if ext in (".xlsx", ".xlsm"):
            parts = []
            for name, rows in sh.sheet_rows(p).items():
                parts.append(f"## {name}")
                parts += [" | ".join(as_text(v) for v in r if not is_blank(v)) for r in rows]
            return "\n".join(parts)[:limit], ""
        if ext == ".pdf":
            try:
                from pypdf import PdfReader  # optional
            except ImportError:
                return "", "PDF text was not extracted (install the optional 'pypdf' package)."
            return "\n".join(pg.extract_text() or "" for pg in PdfReader(str(p)).pages)[:limit], ""
    except (OSError, KeyError, zipfile.BadZipFile, sh.SheetError) as e:
        return "", f"Could not read text: {e}"
    return "", f"Text extraction is not supported for {ext or 'this'} files."


# -- bundle for one ticket ------------------------------------------------------------
PARSERS = {"categories": parse_categories, "dropdowns": parse_dropdowns, "mapping": parse_mapping,
           "ns_master": parse_ns_master, "review_sheet": parse_review_sheet}


class Refs:
    """The references a ticket is processed with.  A file uploaded to the ticket
    overrides the library's active version of the same kind."""

    def __init__(self, sources: dict[str, list[dict]]):
        # sources: kind -> [{"path": Path, "name": str, "origin": "library"|"ticket", "id": ...}]
        self.sources = sources
        self.errors: dict[str, str] = {}
        self.categories, self.dropdowns, self.mapping = [], {}, None
        self.templates: dict[str, dict] = {}
        self.ns_master = None
        self.ns_master_pending = None  # progress while a large master is read in the background
        self.reviews: list[dict] = []
        self.sop_text, self.sop_notes = "", []
        self.problems: list[str] = []
        self._load()

    @classmethod
    def for_ticket(cls, ticket) -> "Refs":
        lib = Library()
        sources = {}
        for kind in REF_KINDS:
            mine = ticket.files(kind)
            if mine:
                pick = mine if kind in MULTI_KINDS else mine[-1:]
                sources[kind] = [{"path": ticket.path(f["stored"]), "name": f["name"], "origin": "ticket",
                                  "id": f["id"], "sha1": f["sha1"]} for f in pick]
            else:
                sources[kind] = [{"path": lib.path(v), "name": v["name"], "origin": "library", "id": v["id"],
                                  "sha1": v["sha1"]} for v in lib.active(kind)]
        return cls(sources)

    @classmethod
    def from_library(cls) -> "Refs":
        lib = Library()
        return cls({k: [{"path": lib.path(v), "name": v["name"], "origin": "library", "id": v["id"],
                         "sha1": v["sha1"]} for v in lib.active(k)] for k in REF_KINDS})

    def first(self, kind):
        s = self.sources.get(kind) or []
        return s[0] if s else None

    def _try(self, kind, fn, *a):
        try:
            return fn(*a)
        except (sh.SheetError, OSError, ValueError, KeyError) as e:
            self.errors[kind] = str(e)
            return None

    def _parsed(self, kind, s, parser, *extra):
        """Parse ``s`` once per file version (see cache.py)."""
        tag = kind + ("-" + "-".join(str(x) for x in extra) if extra else "")
        return self._try(kind, cache.get, tag, s["sha1"], s["path"], lambda p: parser(p, *extra))

    def _load(self):
        if (s := self.first("categories")):
            r = self._parsed("categories", s, parse_categories)
            if r:
                self.categories = r["categories"]
                self.problems += [f"Categories: {p}" for p in r["problems"]]
                self.category_columns = r["columns"]
        if (s := self.first("dropdowns")):
            r = self._parsed("dropdowns", s, parse_dropdowns)
            if r:
                self.dropdowns = r["attributes"]
                self.problems += [f"Dropdowns: {p}" for p in r["problems"]]
        if (s := self.first("mapping")):
            self.mapping = self._parsed("mapping", s, parse_mapping)
            if self.mapping:
                self.problems += [f"Mapping workbook: {p}" for p in self.mapping["problems"]]
        settings = (self.mapping or {}).get("settings", {})
        for kind, prefix in (("m2_template", "m2"), ("ns_template", "ns")):
            if (s := self.first(kind)):
                hr = settings.get(f"{prefix}headerrow")
                t = self._parsed(kind, s, parse_template, settings.get(f"{prefix}templatesheet", ""),
                                 int(float(hr)) if hr and _num(hr) else None)
                if t:
                    self.templates[prefix] = dict(t, path=str(s["path"]))
        if (s := self.first("ns_master")):
            big = Path(s["path"]).stat().st_size > cache.BACKGROUND_BYTES if Path(s["path"]).exists() else False
            if big and not cache.ready("ns_master", s["sha1"]):
                job = cache.background("ns_master", s["sha1"], s["path"], parse_ns_master)
                if job["state"] == "error":
                    self.errors["ns_master"] = job["error"]
                else:
                    self.ns_master_pending = dict(job, name=s["name"])
            else:
                self.ns_master = self._parsed("ns_master", s, parse_ns_master)
        for s in self.sources.get("review_sheet", []):
            r = self._parsed("review_sheet", s, parse_review_sheet)
            if r:
                self.reviews += [dict(rec, file=s["name"]) for rec in r["records"]]
                self.problems += [f"Review sheet {s['name']}: {p}" for p in r["problems"]]
        texts = []
        for s in self.sources.get("sop", []):
            text, note = self._parsed("sop", s, lambda p: list(extract_text(p))) or ("", "")
            if note:
                self.sop_notes.append(f"{s['name']}: {note}")
            if text:
                texts.append(f"# {s['name']}\n{text}")
        self.sop_text = "\n\n".join(texts)

    # -- convenience -------------------------------------------------------------------
    @property
    def content_rules(self) -> list[dict]:
        return (self.mapping or {}).get("content_rules", [])

    def content_rule(self, key: str, category: str = "") -> str:
        """The most specific content rule value for a category (prefix match), else the general one."""
        best, best_len = "", -1
        for r in self.content_rules:
            if r["rule"] != key:
                continue
            c = r["category"]
            if not c and best_len < 0:
                best, best_len = r["value"], 0
            elif c and category and (category == c or category.startswith(c + " >")) and len(c) > best_len:
                best, best_len = r["value"], len(c)
        return best

    def approved_values(self, list_name: str) -> list[str] | None:
        if not list_name:
            return None
        if list_name in self.dropdowns:
            return self.dropdowns[list_name]["values"]
        k = norm(list_name)
        for name, e in self.dropdowns.items():
            if norm(name) == k:
                return e["values"]
        return None

    def option_code(self, list_name: str, value: str) -> str:
        for name, e in self.dropdowns.items():
            if norm(name) == norm(list_name):
                return e["codes"].get(value, "")
        return ""

    def category(self, key: str) -> dict | None:
        return next((c for c in self.categories if c["key"] == key), None)

    def used(self) -> dict:
        return {k: [{"name": s["name"], "origin": s["origin"], "id": s["id"], "sha1": s["sha1"]} for s in v]
                for k, v in self.sources.items() if v}

    def missing(self) -> list[str]:
        return [k for k in REQUIRED_KINDS if not self.sources.get(k)]


def setup_check(refs: Refs) -> list[dict]:
    """Conflicts and gaps between the reference files: [{"severity", "area", "detail"}]."""
    out = []

    def add(sev, area, detail):
        out.append({"severity": sev, "area": area, "detail": detail})

    for k in refs.missing():
        add("blocker", REF_KINDS[k], "Not uploaded. Processing can start, but outputs that depend on it are blocked.")
    for k, e in refs.errors.items():
        add("blocker", REF_KINDS[k], f"Could not be read: {e}")
    for p in refs.problems:
        add("warning", "Reference files", p)
    for n in refs.sop_notes:
        add("warning", "SOP", n)
    if refs.sop_text:
        add("info", "SOP", "SOP text is kept for reference (and as context for the optional assistant). Rules "
            "in it are not applied automatically: write them as Mappings / Attributes / ContentRules rows.")
    m = refs.mapping
    if not m:
        return out
    from .fields import DERIVED_FIELDS  # local import keeps module order simple
    known = set(SOURCE_FIELDS) | set(DERIVED_FIELDS)
    attrs = {a["attribute"] for a in m["attributes"]}
    for system in ("m2", "ns"):
        t = refs.templates.get(system)
        rows = [x for x in m["mappings"] if x["system"] == system]
        if not t:
            continue
        names = [c["name"] for c in t["columns"]]
        mapped = {x["target"] for x in rows}
        label = "Magento 2" if system == "m2" else "NetSuite"
        for x in rows:
            if x["target"] not in names:
                add("blocker", f"{label} mapping", f"Mappings row {x['row']}: '{x['target']}' is not a column of "
                    f"the {label} template.")
            if x["rule"] and x["rule"] not in MAPPING_RULES:
                add("blocker", f"{label} mapping", f"Mappings row {x['row']}: unknown rule '{x['rule']}'.")
            for src in [s for s in re.split(r"[|]", x["source"]) if s.strip()] if x["rule"] in (
                    "copy", "concat", "lookup", "attribute_code") else []:
                f = resolve_field_name(src)
                if not f or (f.startswith("attr:") and f[5:] not in attrs) or (not f.startswith("attr:") and f not in known):
                    add("blocker", f"{label} mapping", f"Mappings row {x['row']}: source '{src}' is not an internal "
                        "field or a defined attribute.")
            if x["list"] and refs.approved_values(x["list"]) is None:
                add("blocker", f"{label} mapping", f"Mappings row {x['row']}: approved list '{x['list']}' is not in "
                    "the dropdown file.")
            if x["rule"] == "lookup" and (x["param"] or x["target"]) not in m["value_maps"]:
                add("blocker", f"{label} mapping", f"Mappings row {x['row']}: value map '{x['param'] or x['target']}'"
                    " does not exist.")
            dv = t["dropdowns"].get(x["target"])
            approved = refs.approved_values(x["list"]) if x["list"] else None
            if dv and approved is not None and set(dv) != set(approved):
                add("warning", f"{label} mapping", f"'{x['target']}': the template drop-down and approved list "
                    f"'{x['list']}' differ ({len(set(dv) - set(approved))} only in template, "
                    f"{len(set(approved) - set(dv))} only in list).")
        for n in names:
            if n not in mapped:
                sev = "blocker" if n in t["required"] else "warning"
                add(sev, f"{label} template", f"Column '{n}' has no mapping row; it will be left empty"
                    + (" but the template marks it required." if sev == "blocker" else "."))
        if t["existing_rows"]:
            add("info", f"{label} template", f"The template sheet already has {t['existing_rows']} data row(s); "
                "the export replaces them with the approved products.")
    for a in m["attributes"]:
        if refs.approved_values(a["list"]) is None:
            add("blocker", "Attributes", f"Attributes row {a['row']}: approved list '{a['list']}' is not in the "
                "dropdown file.")
        if a["source"] and resolve_field_name(a["source"]) is None:
            add("blocker", "Attributes", f"Attributes row {a['row']}: source field '{a['source']}' is unknown.")
        for c in a["categories"]:
            if refs.categories and not any(k["key"] == c or k["key"].startswith(c + " >") for k in refs.categories):
                add("warning", "Attributes", f"Attributes row {a['row']}: category '{c}' is not in the category list.")
    for name, vm in m["value_maps"].items():
        attr = next((a for a in m["attributes"] if a["attribute"] == name), None)
        approved = refs.approved_values(attr["list"] if attr else name)
        if approved is None:
            continue
        for e in vm.values():
            if e["target"] and e["target"] not in approved:
                add("blocker", "Value maps", f"ValueMap '{name}' row {e['row']}: '{e['target']}' is not an approved "
                    "value.")
    if not any(r["rule"] == "title_pattern" for r in m["content_rules"]):
        add("warning", "Content rules", "No title_pattern rule: titles are proposed from the vendor product name "
            "and every title needs manual review.")
    return out


def starter_mapping_workbook(path: Path, refs: Refs) -> None:
    """A Mapping & content rules workbook pre-listing every template column,
    with mapping cells left for the team to fill in."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Settings"
    ws.append(["Setting", "Value", "Notes"])
    for k, note in (("M2 Template Sheet", "Sheet of the Magento 2 template to fill (blank = detect)"),
                    ("M2 Header Row", "Header row number (blank = detect)"),
                    ("NS Template Sheet", "Sheet of the NetSuite template to fill (blank = detect)"),
                    ("NS Header Row", "Header row number (blank = detect)")):
        t = refs.templates.get(k[:2].lower())
        val = (t["sheet"] if "Sheet" in k else t["header_row"]) if t else ""
        ws.append([k, val, note])
    ws = wb.create_sheet("Mappings")
    ws.append(["Target System", "Target Field", "Internal Field", "Rule", "Rule Parameter", "Required",
               "Data Type", "Approved Values List", "Validation Rule", "Max Length", "Multi Value Separator",
               "Source or Approval Reference"])
    for system in ("m2", "ns"):
        t = refs.templates.get(system)
        for c in (t["columns"] if t else []):
            ws.append([system.upper(), c["name"], "", "", "", "Y" if c["name"] in t["required"] else "", "",
                       "", "", "", "", ""])
    ws = wb.create_sheet("Attributes")
    ws.append(["Attribute", "Internal Field", "Applies To Categories", "Mandatory", "Approved Values List",
               "Multi Value Separator", "Expected Format"])
    ws = wb.create_sheet("ContentRules")
    ws.append(["Rule", "Value", "Applies To Category", "Notes"])
    for k, note in CONTENT_RULE_KEYS.items():
        ws.append([k, "", "", note])
    ws = wb.create_sheet("ValueMap")
    ws.append(["Map Name", "Source Value", "Target Value"])
    ws = wb.create_sheet("Help")
    ws.append(["Mapping rule", "Meaning"])
    for k, v in MAPPING_RULES.items():
        ws.append([k, v])
    ws.append([])
    ws.append(["Internal fields", ", ".join(list(SOURCE_FIELDS) + ["category", "category_id", "title",
                                                                   "short_description", "long_description",
                                                                   "bullet_points", "search_terms", "tags"])])
    ws.append(["Attributes", "Use attr:<Attribute name> as Internal Field to output a mapped attribute."])
    for s in wb.worksheets:
        for cell in s[1]:
            cell.font = Font(bold=True)
    wb.save(path)
