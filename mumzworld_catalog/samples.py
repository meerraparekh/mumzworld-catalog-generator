"""Write a small, made-up set of reference and vendor files for trying the app.

    python -m mumzworld_catalog.samples <folder>

Everything here is illustrative test data - it is NOT Mumzworld's real category
list, dropdowns, mappings or templates, and must not be used for production
tickets.  The tests use the same files.
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation


def _wb(path: Path, sheets: dict[str, list[list]]):
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        first = False
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def _png(color: tuple[int, int, int]) -> bytes:
    import struct
    import zlib
    raw = b"".join(b"\x00" + bytes(color) * 4 for _ in range(4))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def write_samples(folder: Path) -> dict[str, Path]:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    out = {}
    out["categories"] = _wb(folder / "sample_categories.xlsx", {"Categories": [
        ["Category ID", "Level 1", "Level 2", "Level 3", "Keywords", "Special"],
        [101, "Baby Gear", "Strollers", "Lightweight Strollers", "buggy|travel stroller|compact stroller", ""],
        [102, "Baby Gear", "Strollers", "Jogging Strollers", "jogger", ""],
        [103, "Baby Gear", "Car Seats", "Infant Car Seats", "car seat|capsule", "Y"],
        [201, "Feeding", "Bottles", "Feeding Bottles", "baby bottle", ""],
        [202, "Feeding", "Bottles", "Bottle Teats", "nipple|teat", ""],
        [301, "Toys", "Soft Toys", "Plush Toys", "teddy|plush|stuffed", ""],
        [302, "Toys", "Learning Toys", "Stacking Toys", "stacker|stacking", ""],
    ]})
    out["dropdowns"] = _wb(folder / "sample_dropdowns.xlsx", {"Lists": [
        ["Attribute", "Value", "Option ID"],
        ["Colour", "Black", "c1"], ["Colour", "Grey", "c2"], ["Colour", "Blue", "c3"], ["Colour", "Pink", "c4"],
        ["Colour", "Multicolour", "c5"], ["Colour", "White", "c6"],
        ["Age Group", "0-6 Months", "a1"], ["Age Group", "6-12 Months", "a2"], ["Age Group", "1-3 Years", "a3"],
        ["Age Group", "3+ Years", "a4"],
        ["Material", "Plastic", "m1"], ["Material", "Polyester", "m2"], ["Material", "Glass", "m3"],
        ["Material", "Silicone", "m4"], ["Material", "Cotton", "m5"],
    ]})
    out["mapping"] = _wb(folder / "sample_mapping_rules.xlsx", {
        "Settings": [["Setting", "Value"], ["M2 Template Sheet", "Upload"], ["M2 Header Row", 1]],
        "Mappings": [
            ["Target System", "Target Field", "Internal Field", "Rule", "Rule Parameter", "Required", "Data Type",
             "Approved Values List", "Validation Rule", "Max Length"],
            ["M2", "sku", "sku", "copy", "", "Y", "text", "", "", 64],
            ["M2", "name", "title", "copy", "", "Y", "text", "", "", 120],
            ["M2", "short_description", "short_description", "copy", "", "Y", "", "", "", ""],
            ["M2", "description", "long_description", "copy", "", "Y", "", "", "", ""],
            ["M2", "categories", "", "category_path", "/", "Y", "", "", "", ""],
            ["M2", "color", "attr:Colour", "copy", "", "Y", "", "Colour", "", ""],
            ["M2", "age_group", "attr:Age Group", "copy", "", "", "", "Age Group", "", ""],
            ["M2", "price", "retail_price", "copy", "", "Y", "decimal", "", "", ""],
            ["M2", "attribute_set_code", "", "constant", "Default", "Y", "", "", "", ""],
            ["M2", "base_image", "", "images", ",", "", "", "", "", ""],
            ["NS", "Item Name/Number", "sku", "copy", "", "Y", "text", "", "", ""],
            ["NS", "Display Name", "", "template", "{title}", "Y", "", "", "", ""],
            ["NS", "UPC Code", "barcode", "copy", "", "Y", "text", "", r"\d{8,14}", ""],
            ["NS", "Class", "", "category_level", "2", "Y", "", "", "", ""],
            ["NS", "Category ID", "", "category_id", "", "Y", "", "", "", ""],
            ["NS", "Brand", "brand", "copy", "", "Y", "", "", "", ""],
            ["NS", "Cost", "cost_price", "copy", "", "", "decimal", "", "", ""],
        ],
        "Attributes": [
            ["Attribute", "Internal Field", "Applies To Categories", "Mandatory", "Approved Values List",
             "Multi Value Separator", "Expected Format"],
            ["Colour", "colour", "", "Y", "Colour", "", ""],
            ["Age Group", "age_group", "Toys|Baby Gear", "N", "Age Group", "", ""],
            ["Material", "material", "Feeding", "Y", "Material", "", ""],
        ],
        "ContentRules": [
            ["Rule", "Value", "Applies To Category"],
            ["title_pattern", "{brand} {product_name} - {colour}", ""],
            ["title_max_length", "80", ""],
            ["title_case", "title", ""],
            ["short_description_max_length", "160", ""],
            ["bullets_max", "5", ""],
            ["images_required", "Y", ""],
        ],
        "ValueMap": [["Map Name", "Source Value", "Target Value"], ["Colour", "Charcoal", "Grey"],
                     ["Colour", "Navy", "Blue"]],
    })
    wb = Workbook()
    ws = wb.active
    ws.title = "Upload"
    ws.append(["sku", "name", "short_description", "description", "categories", "color", "age_group", "price",
               "attribute_set_code", "base_image"])
    ws.append(["EXAMPLE-1", "Example row", "", "", "", "", "", "", "", ""])
    dv = DataValidation(type="list", formula1='"Black,Grey,Blue,Pink,Multicolour,White"')
    dv.add("F2:F2000")
    ws.add_data_validation(dv)
    wb.create_sheet("Instructions").append(["Do not change the header row."])
    wb.save(folder / "sample_m2_template.xlsx")
    out["m2_template"] = folder / "sample_m2_template.xlsx"
    (folder / "sample_ns_template.csv").write_text(
        "Item Name/Number,Display Name,UPC Code,Class,Category ID,Brand,Cost\n", "utf-8")
    out["ns_template"] = folder / "sample_ns_template.csv"
    out["ns_master"] = _wb(folder / "sample_netsuite_master.xlsx", {"Items": [
        ["Internal ID", "Item Name/Number", "UPC Code", "Display Name", "Brand"],
        [9001, "TOT-100", "5012345678900", "Totsy Old Stroller", "Totsy"],
        [9002, "OLD-555", "4006381333931", "Other Brand Bottle", "Other"],
    ]})
    out["review_sheet"] = _wb(folder / "sample_review_sheet.xlsx", {"Reviews": [
        ["SKU", "Field", "Original Value", "Corrected Value", "Reviewer Comments", "Approval Status", "Category",
         "Brand", "Date"],
        ["X1", "Colour", "Gray", "Grey", "British spelling", "Approved", "Toys > Soft Toys > Plush Toys", "Cuddle", "2026-05-01"],
        ["X2", "Colour", "Gray", "Grey", "British spelling", "Approved", "Feeding > Bottles > Feeding Bottles",
         "Sippy", "2026-05-03"],
        ["X3", "Title", "cuddle bear", "Cuddle Bear Plush Toy - Pink", "Add type", "Approved", "", "Cuddle", ""],
    ]})
    (folder / "sample_sop.txt").write_text("Sample SOP (illustrative).\nTitles: Brand + product name + colour.\n",
                                           "utf-8")
    out["sop"] = folder / "sample_sop.txt"
    out["vendor_input"] = _wb(folder / "sample_vendor_input.xlsx", {"Products": [
        ["Vendor SKU", "Brand", "Product Name", "EAN", "Colour", "Category", "Age", "Material", "Description",
         "Key Features", "RRP", "Cost"],
        ["TOT-200", "Totsy", "Totsy Breeze Compact Stroller", "5012345678917", "Charcoal", "Strollers", "0-6 Months",
         "Polyester", "A compact stroller for travel. Folds with one hand.\n\nSuitable from birth.",
         "One-hand fold\nCarry strap\nUV canopy", 499, 300],
        ["TOT-100", "Totsy", "Totsy Old Stroller", "5012345678900", "Black", "Strollers", "", "", "Old model.", "",
         399, 200],
        ["SIP-01", "Sippy", "Glass Baby Bottle 240ml", "4006381333931", "Clear", "Bottles", "0-6 Months", "",
         "A glass feeding bottle. BPA free.", "Anti-colic vent", 59, 30],
        ["CUD-7", "Cuddle", "Cuddle Bear", "", "Pink", "", "1-3 Years", "Cotton",
         "A soft teddy bear for cuddles.", "Soft fabric", 89, "40"],
        ["CUD-7", "Cuddle", "Cuddle Bear Big", "12345", "Pink", "Plush", "", "", "", "", 120, 60],
    ]})
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("images/TOT-200_1.png", _png((40, 40, 40)))
        zf.writestr("images/TOT-200_2.png", _png((60, 60, 60)))
        zf.writestr("images/SIP-01.png", _png((240, 240, 255)))
        zf.writestr("images/xxCUD-7yy.png", _png((255, 180, 200)))
        zf.writestr("images/TOT-200alt.png", _png((61, 60, 60)))
        zf.writestr("__MACOSX/._TOT-200_1.png", b"junk")
        zf.writestr("../evil.png", _png((0, 0, 0)))
    (folder / "sample_images.zip").write_bytes(buf.getvalue())
    out["images"] = folder / "sample_images.zip"
    out["response"] = _wb(folder / "sample_vendor_response.xlsx", {"Answers": [
        ["SKU", "Material", "Barcode", "Colour"],
        ["SIP-01", "Glass", "", "White"],
        ["UNKNOWN-1", "Plastic", "", ""],
    ]})
    return out


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "mumzworld_samples")
    for k, p in write_samples(target).items():
        print(f"{k:14} {p}")
